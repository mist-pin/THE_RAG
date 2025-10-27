"""
Text Utilities Module

This module provides functions for extracting text from various document formats
(PDF, DOCX, TXT, MD, CSV, etc.) and counting tokens using tiktoken.
"""

import logging
from pathlib import Path
from typing import Dict, Any
import tiktoken
from PyPDF2 import PdfReader
from docx import Document
import csv


# Configure logging
logger = logging.getLogger(__name__)

# Supported file types
SUPPORTED_EXTENSIONS = ['.pdf', '.docx', '.txt', '.md', '.csv', '.log', '.json', '.xml', '.html']


def get_file_info(file_path: str) -> Dict[str, Any]:
    """
    Get information about a file including type and size.
    
    Args:
        file_path (str): Path to the file
        
    Returns:
        Dict[str, Any]: Dictionary with file information
    """
    path = Path(file_path)
    
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    
    file_size_bytes = path.stat().st_size
    file_size_kb = file_size_bytes / 1024
    file_size_mb = file_size_kb / 1024
    
    extension = path.suffix.lower()
    
    # Determine file type category
    file_type = "Unknown"
    if extension == '.pdf':
        file_type = "PDF Document"
    elif extension == '.docx':
        file_type = "Word Document"
    elif extension in ['.txt', '.log']:
        file_type = "Text File"
    elif extension == '.md':
        file_type = "Markdown File"
    elif extension == '.csv':
        file_type = "CSV File"
    elif extension == '.json':
        file_type = "JSON File"
    elif extension in ['.xml', '.html']:
        file_type = "Markup File"
    
    info = {
        "file_path": str(path.absolute()),
        "file_name": path.name,
        "file_type": file_type,
        "extension": extension,
        "size_bytes": file_size_bytes,
        "size_kb": round(file_size_kb, 2),
        "size_mb": round(file_size_mb, 2),
    }
    
    logger.info(f"File info: {info['file_name']} ({info['file_type']}, {info['size_kb']} KB)")
    
    return info


def extract_text(file_path: str) -> str:
    """
    Extract text from a document file.
    
    Supports PDF, DOCX, TXT, MD, CSV, JSON, XML, HTML, and LOG file formats.
    
    Args:
        file_path (str): Path to the document file
        
    Returns:
        str: Extracted text content
        
    Raises:
        FileNotFoundError: If the file doesn't exist
        ValueError: If the file format is not supported
    """
    logger.info(f"Extracting text from: {file_path}")
    
    path = Path(file_path)
    
    # Check if file exists
    if not path.exists():
        logger.error(f"File not found: {file_path}")
        raise FileNotFoundError(f"File not found: {file_path}")
    
    # Get file extension
    extension = path.suffix.lower()
    
    try:
        # Extract based on file type
        if extension == '.pdf':
            text = _extract_from_pdf(file_path)
        elif extension == '.docx':
            text = _extract_from_docx(file_path)
        elif extension in ['.txt', '.md', '.log', '.json', '.xml', '.html']:
            text = _extract_from_txt(file_path)
        elif extension == '.csv':
            text = _extract_from_csv(file_path)
        else:
            logger.error(f"Unsupported file format: {extension}")
            raise ValueError(f"Unsupported file format: {extension}. Supported formats: {', '.join(SUPPORTED_EXTENSIONS)}")
        
        logger.info(f"Successfully extracted {len(text)} characters from {file_path}")
        return text
        
    except Exception as e:
        logger.error(f"Error extracting text from {file_path}: {str(e)}")
        raise


def _extract_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file.
    
    Args:
        file_path (str): Path to the PDF file
        
    Returns:
        str: Extracted text content
    """
    logger.debug(f"Extracting text from PDF: {file_path}")
    
    text = ""
    with open(file_path, 'rb') as file:
        pdf_reader = PdfReader(file)
        num_pages = len(pdf_reader.pages)
        logger.debug(f"PDF has {num_pages} pages")
        
        for page_num, page in enumerate(pdf_reader.pages, 1):
            page_text = page.extract_text()
            text += page_text
            logger.debug(f"Extracted {len(page_text)} characters from page {page_num}")
    
    return text


def _extract_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file.
    
    Args:
        file_path (str): Path to the DOCX file
        
    Returns:
        str: Extracted text content
    """
    logger.debug(f"Extracting text from DOCX: {file_path}")
    
    doc = Document(file_path)
    text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    
    logger.debug(f"Extracted text from {len(doc.paragraphs)} paragraphs")
    return text


def _extract_from_txt(file_path: str) -> str:
    """
    Extract text from a TXT file.
    
    Args:
        file_path (str): Path to the TXT file
        
    Returns:
        str: Extracted text content
    """
    logger.debug(f"Extracting text from TXT: {file_path}")
    
    with open(file_path, 'r', encoding='utf-8') as file:
        text = file.read()
    
    return text


def _extract_from_csv(file_path: str) -> str:
    """
    Extract text from a CSV file.
    
    Args:
        file_path (str): Path to the CSV file
        
    Returns:
        str: Extracted text content (formatted as readable text)
    """
    logger.debug(f"Extracting text from CSV: {file_path}")
    
    text_parts = []
    
    with open(file_path, 'r', encoding='utf-8') as file:
        csv_reader = csv.reader(file)
        
        # Read header
        try:
            header = next(csv_reader)
            text_parts.append("Headers: " + ", ".join(header))
            text_parts.append("")  # Empty line
        except StopIteration:
            pass
        
        # Read rows
        for row_num, row in enumerate(csv_reader, 1):
            row_text = " | ".join(row)
            text_parts.append(f"Row {row_num}: {row_text}")
    
    return "\n".join(text_parts)


def count_tokens(text: str, model: str = "gpt-3.5-turbo") -> int:
    """
    Count the number of tokens in a text string using tiktoken.
    
    Args:
        text (str): The text to count tokens for
        model (str): The model name for tokenization (default: "gpt-3.5-turbo")
        
    Returns:
        int: Number of tokens in the text
    """
    logger.debug(f"Counting tokens for text of length {len(text)} characters")
    
    try:
        # Get the encoding for the specified model
        encoding = tiktoken.encoding_for_model(model)
        
        # Encode the text and count tokens
        tokens = encoding.encode(text)
        token_count = len(tokens)
        
        logger.info(f"Token count: {token_count} tokens for {len(text)} characters")
        return token_count
        
    except Exception as e:
        logger.error(f"Error counting tokens: {str(e)}")
        raise


def get_text_excerpt(text: str, max_chars: int = 1000) -> str:
    """
    Get an excerpt from the beginning of the text.
    
    Args:
        text (str): The full text
        max_chars (int): Maximum number of characters to extract
        
    Returns:
        str: Text excerpt
    """
    logger.debug(f"Extracting excerpt of up to {max_chars} characters")
    
    excerpt = text[:max_chars]
    
    if len(text) > max_chars:
        logger.info(f"Extracted excerpt: {len(excerpt)} characters (truncated from {len(text)})")
    else:
        logger.info(f"Extracted full text as excerpt: {len(excerpt)} characters")
    
    return excerpt
